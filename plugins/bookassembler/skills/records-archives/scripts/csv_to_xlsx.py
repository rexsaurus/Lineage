#!/usr/bin/env python3
"""Build a formatted .xlsx workbook from one or more CSV files (one sheet each).

Usage: csv_to_xlsx.py out.xlsx "Sheet Name=path/to/file.csv" ["Other=other.csv" ...]
CSV stays the source of truth (easy for agents to diff and edit); the workbook is the
human-friendly view. Re-run after every CSV change.
"""
import csv, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="2F3B4C")
FLAG_FILLS = {  # colour whole row by a status/confidence column if present
    "confirmed": "E3F1E0", "high": "E3F1E0",
    "inferred": "FFF4D6", "medium": "FFF4D6", "estimated": "FFF4D6",
    "unverified": "FBE2E2", "low": "FBE2E2", "conflict": "FBE2E2", "lost": "EDEDED",
}
STATUS_COLS = ("status", "confidence", "date_confidence")


def add_sheet(wb, name, path):
    rows = list(csv.reader(open(path, newline="", encoding="utf-8")))
    ws = wb.create_sheet(name[:31])
    if not rows:
        return
    for r in rows:
        ws.append(r)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = HEADER_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
    header = [h.strip().lower() for h in rows[0]]
    status_idx = next((header.index(s) for s in STATUS_COLS if s in header), None)
    for i, col in enumerate(zip(*rows), start=1):
        longest = max(len(str(v)) for v in col)
        ws.column_dimensions[get_column_letter(i)].width = min(max(10, longest + 2), 60)
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
        if status_idx is not None:
            key = str(row[status_idx].value or "").strip().lower()
            fill = FLAG_FILLS.get(key)
            if fill:
                for c in row:
                    c.fill = PatternFill("solid", fgColor=fill)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    out = Path(sys.argv[1])
    wb = Workbook()
    wb.remove(wb.active)
    for spec in sys.argv[2:]:
        name, path = spec.split("=", 1)
        add_sheet(wb, name, path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"wrote {out}")


if __name__ == "__main__":
    main()
